#!/usr/bin/env python3
"""Extract selected sheets from the Wallen Excel workbook and save as TSVs.

This script reads data/raw/PRJNA834801/source_data/Source_Data_24Oct2022.xlsx
and extracts the sheets: subject_metadata, metaphlan_rel_ab, and
humann_pathway_counts. Each sheet is written as a TSV to
data/processed/wallen/. The script prints the dimensions of each extracted
table.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Dict

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
INPUT_XLSX = (
    ROOT / "data" / "raw" / "PRJNA834801" / "source_data" / "Source_Data_24Oct2022.xlsx"
)
OUT_DIR = ROOT / "data" / "processed" / "wallen"
OUT_DIR.mkdir(parents=True, exist_ok=True)

TARGET_SHEETS = [
    "subject_metadata",
    "metaphlan_rel_ab",
    "humann_pathway_counts",
]


def normalize_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (value or "").lower())


def find_sheet_map(workbook: openpyxl.workbook.workbook.Workbook) -> Dict[str, str]:
    """Return a mapping from normalized sheet name -> actual sheet title."""
    mapping: Dict[str, str] = {}
    for name in workbook.sheetnames:
        mapping[normalize_name(name)] = name
    return mapping


def extract_sheet_to_tsv(workbook: openpyxl.workbook.workbook.Workbook, actual_name: str, outpath: Path) -> tuple[int, int]:
    ws = workbook[actual_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        # empty sheet
        with outpath.open("w", encoding="utf-8", newline="") as fh:
            pass
        return 0, 0

    header = [str(c) if c is not None else "" for c in rows[0]]
    data_rows = rows[1:]

    with outpath.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter="\t")
        writer.writerow(header)
        for r in data_rows:
            # normalize cell values to strings (empty for None)
            writer.writerow(["" if v is None else str(v) for v in r])

    return len(data_rows), len(header)


def main() -> None:
    if not INPUT_XLSX.exists():
        print(f"Input workbook not found: {INPUT_XLSX}")
        return

    wb = openpyxl.load_workbook(INPUT_XLSX, read_only=True, data_only=True)
    sheet_map = find_sheet_map(wb)

    for target in TARGET_SHEETS:
        norm = normalize_name(target)
        if norm not in sheet_map:
            print(f"Sheet '{target}' not found in workbook (tried normalized name '{norm}'). Skipping.")
            continue

        actual = sheet_map[norm]
        outpath = OUT_DIR / f"{target}.tsv"
        rows, cols = extract_sheet_to_tsv(wb, actual, outpath)
        print(f"{target}: {rows} rows x {cols} columns -> {outpath.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
